"""Read FIX workbooks; retain dialogue text without rewriting."""
import openpyxl,json,re,pathlib
root=pathlib.Path(__file__).resolve().parents[1]
base=pathlib.Path('/workspace/scratch/702c6dc06a78')
w=openpyxl.load_workbook(base/'upload/BAR_ECLUNE_CASE01-06_全体収録台本_完成版(1).xlsx',data_only=True)
s=w['01_CASE01全体台本'];rows={}
for i,r in enumerate(s.iter_rows(values_only=True),1):
 if i==1:continue
 ref=re.search(r'参照ID：([^\s]+)',str(r[7] or ''))
 rows[i]={'row':i,'id':ref[1] if ref else f'C01-ROW-{i}','no':r[0],'section':r[1],'kind':r[2],'speaker':r[3] or '', 'text':r[4] or '', 'note':r[5] or ''}
for r in rows.values():
 if r['text']:
  r['text']=r['text'].replace('\\u3000','\u3000').replace('\\n','\n') # Excel literal escape of fullwidth space/newline

def seq(a,b):return [rows[i] for i in range(a,b+1) if rows[i]['kind'] in ['台詞','地の文','見出し'] and rows[i]['text'] and '必須未完了' not in rows[i]['note'] and '必須完了後' not in rows[i]['note']]
def point(key,label,a,b,loc='店内',requires=None,ev=None,images=None,typ='object',optional=False):
 return dict(key=key,label=label,rows=seq(a,b),location=loc,requires=requires or [],evidence=ev or [],images=images or [],type=typ,optional=optional)
routes={}
routes['紅']=dict(character='紅',intro=seq(156,166),exit=seq(282,310),blocked=rows[280],ready=rows[281],points=[
 point('register','レジ',171,181,ev=['E01'],images=['112'],typ='record'),point('interior','店内',183,188),point('owner','宮本',190,202,typ='person'),point('desk','机',206,212,'事務室'),point('envelope','封筒',214,224,'事務室',ev=['E02'],images=['113']),point('window','窓',226,234,'事務室'),point('door','裏口のドア',236,243,'事務室',ev=['E03'],requires=['window']),point('bag','鞄',246,279,'事務室',requires=['window','door'],ev=['E04','E11'],images=['114','115'],optional=True)],requiredEvidence=['E01','E02','E03'])
routes['藍']=dict(character='藍',intro=seq(311,335),exit=seq(463,505),blocked=rows[461],ready=rows[462],points=[
 point('register','レジ',339,348,ev=['E01'],images=['112'],typ='record'),point('camera','防犯カメラ',350,359,typ='record'),point('cam1','CAM 01',362,363,'防犯カメラ',requires=['camera'],images=['116'],typ='record'),point('cam2','CAM 02',365,366,'防犯カメラ',requires=['camera'],images=['117'],typ='record'),point('cam3','CAM 03',368,369,'防犯カメラ',requires=['camera'],images=['118'],typ='record'),point('cam4','CAM 04',371,381,'防犯カメラ',requires=['camera','cam1','cam2','cam3'],ev=['E05'],images=['119-01'],typ='record'),point('door','裏口のドア',385,394,'事務室'),point('window','窓',396,414,'事務室',requires=['door'],ev=['E06']),point('cam4-new','CAM 04：追加確認',417,460,'防犯カメラ',requires=['cam4','door','window'],ev=['E07'],images=['119-02','119-03'],typ='record')],requiredEvidence=['E01','E05','E06','E07'])
routes['翠']=dict(character='翠',intro=seq(506,528),exit=seq(745,801),blocked=rows[743],ready=rows[744],points=[
 point('register','レジ',532,535,ev=['E01'],images=['112'],typ='record'),point('camera','防犯カメラ',537,541,ev=['E05'],typ='record',optional=True),point('envelope','封筒',543,551,ev=['E02'],images=['113'],optional=True),point('kobayashi','小林美月',553,587,typ='person',optional=True),point('tanabe','田辺翔太',589,620,ev=['E10'],requires=['kobayashi'],typ='person',optional=True),point('owner','宮本浩司',622,675,ev=['E08'],typ='person'),point('brother','宮本俊介',678,742,requires=['owner'],ev=['E09','E11'],typ='person')],requiredEvidence=['E01','E08','E09'])
for key,a,b,ia,ib,block,ready,pts,required in [
 ('紅→藍',803,835,803,817,833,834,[point('camera','防犯カメラ',821,832,ev=['E05','E06','E07'],images=['119-02','119-03'],typ='record')],['E07']),
 ('紅→翠',836,896,836,850,894,895,[point('owner','宮本浩司',854,860,ev=['E08','E11'],typ='person'),point('tanabe','田辺翔太',862,867,ev=['E10'],typ='person'),point('brother','宮本俊介',870,893,requires=['owner','tanabe'],ev=['E09'],typ='person')],['E08','E09','E11']),
 ('藍→紅',897,927,897,907,925,926,[point('bag','鞄',911,921,ev=['E04'],images=['114','115']),point('owner','宮本浩司',924,924,requires=['bag'],ev=['E11'],typ='person')],['E04']),
 ('藍→翠',928,980,928,937,978,979,[point('owner','宮本浩司',941,946,ev=['E08','E11'],typ='person'),point('tanabe','田辺翔太',948,953,ev=['E10'],typ='person'),point('brother','宮本俊介',956,977,requires=['owner','tanabe'],ev=['E09'],typ='person')],['E08','E09','E11']),
 ('翠→紅',981,1016,981,991,1014,1015,[point('bag','鞄',995,999,ev=['E04'],images=['114','115']),point('brother','宮本俊介',1002,1013,requires=['bag'],typ='person')],['E04']),
 ('翠→藍',1017,1047,1017,1028,1045,1046,[point('camera','防犯カメラ',1032,1044,ev=['E05','E06','E07'],images=['119-02','119-03'],typ='record')],['E07'])]:
 routes[key]=dict(character=key[-1],intro=seq(ia,ib),exit=[],blocked=rows[block],ready=rows[ready],points=pts,requiredEvidence=required)
# Evidence descriptions from authoritative specification, not story rewriting.
spec=openpyxl.load_workbook(base/'game-inputs/BAR_ECLUNE_ゲーム仕様書_v6.2_CASE04時刻ロジック_FIX(1).xlsx',data_only=True)
ev={}
for r in spec['09_CASE01確定仕様'].iter_rows(values_only=True):
 if isinstance(r[0],str) and re.fullmatch(r'E\d\d',r[0]):ev[r[0]]={'id':r[0],'category':r[1],'name':r[2],'description':r[4],'images':[]}
for k,img in {'E01':['112'],'E02':['113'],'E04':['114','115'],'E05':['116','117','118','119-01'],'E07':['119-02','119-03']}.items():ev[k]['images']=img
revisit={}
for r in spec['23_再調査時セリフ'].iter_rows(values_only=True):
 if r[0] in ['R01','R02','R03','R04']:revisit[r[0]]=dict(zip(['紅','藍','翠'],r[3:6]))
questions=[]
old=spec['11_CASE01確定台本']
for textrow,optionrows,correct in [(1497,[1498,1499],1),(1503,[1504,1505,1506,1507],1),(1514,[1515,1516,1517,1518],2)]:
 questions.append({'text':old.cell(textrow,6).value,'options':[re.sub(r'^[A-D] ','',old.cell(n,6).value) for n in optionrows],'correct':correct})
allrows=list(rows.values())
data={'title':'BAR ECLUNE','case':1,'name':'消えた売上金','rows':allrows,'prologue':seq(2,155),'routes':routes,'day2End':seq(1048,1059),'day3Intro':seq(1060,1082),'questions':questions,'rescue':seq(1090,1090),'evidence':ev,'revisit':revisit,'verdict':{'perfect':seq(1095,1108),'goodIntro':seq(1110,1113),'goodQuestions':[ [rows[i] for i in range(a,b+1) if rows[i]['kind'] in ['台詞','地の文','UI']] for a,b in [(1115,1120),(1122,1127),(1129,1135)] ],'goodEnd':seq(1136,1137),'badIntro':seq(1139,1146),'badQuestions':[[rows[i] for i in range(a,b+1) if rows[i]['kind'] in ['台詞','地の文','UI']] for a,b in [(1148,1154),(1156,1160),(1162,1168)]],'badEnd':seq(1169,1169),'cash':seq(1170,1172),'truth':seq(1173,1224)}}
(root/'data.js').write_text('window.ECLUNE_DATA = '+json.dumps(data,ensure_ascii=False,separators=(',',':'))+';\n',encoding='utf-8')
(root/'audio-manifest.js').write_text('window.ECLUNE_AUDIO = '+json.dumps({'bgm':{str(i):f'audio/bgm/bgm-{i:03}.wav' for i in [2,5,7,8,9]},'se':{str(i):f'audio/se/se-{i:03}.mp3' for i in [10,11,12,13,14,15]},'silentEvents':['ui','evidence','return-unlock','grade','clear','text']},ensure_ascii=False)+';\n')
(root/'voice-manifest.js').write_text('window.ECLUNE_VOICES = {}; // Map FIX reference IDs to relative voice files after CV delivery.\n')
(root/'tools/source-audit.json').write_text(json.dumps({'source':str(s.title),'rowCount':len(rows),'routing':'manual boundaries matched against workbook','provisionalReturnConditions':{k:v['requiredEvidence'] for k,v in routes.items()},'notes':['完成版C01-GRADEを優先：3問正解PERFECT、1〜2問GOOD、0問BAD。旧仕様書の1問BADとは異なる。','CASE01のルート別必須ID表が旧仕様書にないため、完成台本の核心と解放順から帰還条件を明示。紅DAY1のE04は旧仕様書どおり任意。','鞄未取得時は紅DAY1終端と紅→藍導入の現金発見を前提にする台詞を取得条件でスキップ。本文は改変しない。','CASE02〜06は後続実装。CASE01クリアでCASE02解放状態を保存するが未実装の本編を起動しない。']},ensure_ascii=False,indent=2))
print('compiled',len(rows),'source rows,',len(routes),'routes')
